"""Batch save does ONE read-modify-write of a workbook per phase, not one per document: spies on
openpyxl.load_workbook and Workbook.save count only calls on the exports-dir workbook."""

from pathlib import Path
from typing import Any

import openpyxl
import pytest
from openpyxl.workbook.workbook import Workbook

from app.features.excel import service as excel_service
from app.features.excel.batch_test_support import (
    BASE,
    CLOCK,
    make_doc,
    physical,
    post_bulk,
    seed_book,
    seed_old_record,
    seed_settings,
    write_files,
)

WB26 = "Book2026"


class Spy:
    def __init__(self, stem: str, d: Path) -> None:
        self.stem, self.d = stem, d
        self.loads = 0
        self.saves = 0

    def mine(self, target: Any) -> bool:
        p = Path(str(target))
        return p.parent == self.d and p.name.startswith(self.stem)


@pytest.fixture
def exports_dir(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    d = tmp_path / "exports"
    d.mkdir()
    monkeypatch.setattr(excel_service, "EXPORT_DIR", d)
    monkeypatch.setattr(excel_service, "now_ist", lambda: CLOCK)
    return d


def _install(monkeypatch: pytest.MonkeyPatch, spy: Spy) -> None:
    real_load = openpyxl.load_workbook
    real_save = Workbook.save

    def load(filename: Any, *a: Any, **kw: Any) -> Any:
        if spy.mine(filename):
            spy.loads += 1
        return real_load(filename, *a, **kw)

    def save(self: Workbook, filename: Any) -> None:
        if spy.mine(filename):
            spy.saves += 1
        real_save(self, filename)

    monkeypatch.setattr(openpyxl, "load_workbook", load)
    monkeypatch.setattr(Workbook, "save", save)


async def test_bulk_of_30_with_10_save_again_loads_and_saves_the_file_at_most_twice(
    client: Any, db: Any, make_user: Any, exports_dir: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    user = await make_user()
    wb = await seed_book(db, user.id, WB26, 2026, True)
    await seed_settings(db, user.id, WB26, 2026)
    old_rows = []
    docs = []
    for i in range(30):
        d = await make_doc(db, user.id, f"E{i:02d}", "02/03/2026")
        docs.append(d)
        if i % 3 == 0:  # 10 documents are saved again
            await seed_old_record(db, user.id, d, wb, WB26, "March", f"E{i:02d}", "02/03/2026")
            old_rows.append(["Delivery Challan", f"E{i:02d}", "02/03/2026", "t0"])
    write_files(exports_dir, {physical(user.id, WB26): {"March": old_rows}})

    spy = Spy(physical(user.id, WB26), exports_dir)
    _install(monkeypatch, spy)
    r = await post_bulk(client, user, docs)
    body = r.json()
    assert r.status_code == 200 and len(body["succeeded"]) == 30 and body["replaced"] == 10
    assert 1 <= spy.loads <= 2, spy.loads
    assert 1 <= spy.saves <= 2, spy.saves


async def test_bulk_of_30_new_documents_loads_and_saves_the_file_once(
    client: Any, db: Any, make_user: Any, exports_dir: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    user = await make_user()
    await seed_book(db, user.id, WB26, 2026, True)
    await seed_settings(db, user.id, WB26, 2026)
    write_files(exports_dir, {physical(user.id, WB26): {"March": []}})
    docs = [await make_doc(db, user.id, f"F{i:02d}", "02/03/2026") for i in range(30)]

    spy = Spy(physical(user.id, WB26), exports_dir)
    _install(monkeypatch, spy)
    body = (await post_bulk(client, user, docs)).json()
    assert len(body["succeeded"]) == 30 and body["replaced"] == 0
    assert (spy.loads, spy.saves) == (1, 1)


async def test_single_save_is_one_load_and_one_save(
    client: Any, db: Any, make_user: Any, exports_dir: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    user = await make_user()
    await seed_book(db, user.id, WB26, 2026, True)
    await seed_settings(db, user.id, WB26, 2026)
    write_files(exports_dir, {physical(user.id, WB26): {"March": []}})
    d = await make_doc(db, user.id, "G1", "02/03/2026")
    spy = Spy(physical(user.id, WB26), exports_dir)
    _install(monkeypatch, spy)
    assert (await client.post(f"{BASE}/{d}/save", headers=user.headers)).status_code == 200
    assert (spy.loads, spy.saves) == (1, 1)
