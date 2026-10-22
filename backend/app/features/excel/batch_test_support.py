"""Shared helpers for the batch-save tests (test_bulk_batch_*.py). Not a test module.

Everything here works on the DB through the `db` fixture's handle and on an exports directory
that the tests redirect with monkeypatch.setattr(excel_service, "EXPORT_DIR", some_tmp_path) -
never the real backend/exports folder."""

import re
from datetime import UTC, datetime, timedelta, timezone
from pathlib import Path
from typing import Any

import openpyxl
from bson import ObjectId

from app.features.excel import old_excel_oracle as oracle
from app.features.excel import service as excel_service

IST = timezone(timedelta(hours=5, minutes=30))
CLOCK = datetime(2026, 3, 10, 10, 0, tzinfo=IST)
BASE = "/api/documents"
COLLECTIONS = ("documents", "settings", "workbooks", "exportedrows", "auditlogs")
_ISO_RE = re.compile(r"^\d{4}-\d{2}-\d{2}T\d{2}:\d{2}:\d{2}")


def physical(user_id: ObjectId, name: str) -> str:
    return excel_service.physical_workbook_filename(user_id, name)


def xlsx_path(d: Path, user_id: ObjectId, name: str) -> Path:
    return d / f"{physical(user_id, name)}.xlsx"


async def make_doc(
    db: Any,
    owner_id: ObjectId,
    number: str | None,
    date: str | None,
    doc_type: str = "Delivery Challan",
    status: str = "processed",
    tax_invoice_no: str | None = None,
    reference_no: str | None = None,
    is_deleted: bool = False,
) -> ObjectId:
    now = datetime.now(UTC)
    res = await db.documents.insert_one(
        {
            "userId": owner_id,
            "autoName": number or tax_invoice_no or "x",
            "originalFilename": f"{number or tax_invoice_no}.pdf",
            "mimeType": "application/pdf",
            "size": 10,
            "uploadStatus": status,
            "documentType": doc_type,
            "taxInvoiceNo": tax_invoice_no,
            "referenceNo": reference_no,
            "number": number,
            "date": date,
            "edited": False,
            "exported": False,
            "isDeleted": is_deleted,
            "filePurged": True,
            "createdAt": now,
            "updatedAt": now,
        }
    )
    return res.inserted_id


def number_cell(doc_type: str, tin: str | None, ref: str | None, number: str | None) -> Any:
    return oracle._format_number_cell(
        {"documentType": doc_type, "taxInvoiceNo": tin, "referenceNo": ref, "number": number}
    )


# files model: {physical stem: {sheet: [row cells...]}} (rows exclude the header)
Files = dict[str, dict[str, list[list[Any]]]]


def write_files(d: Path, files: Files) -> None:
    d.mkdir(parents=True, exist_ok=True)
    for stem, sheets in files.items():
        book = openpyxl.Workbook()
        book.remove(book.active)
        for sheet_name, rows in sheets.items():
            ws = book.create_sheet(sheet_name)
            ws.append(excel_service.HEADERS)
            for cell in ws[1]:
                cell.font = openpyxl.styles.Font(bold=True)
            for row in rows:
                ws.append(row)
        book.save(d / f"{stem}.xlsx")


def read_dir(d: Path, normalize_timestamps: bool = True) -> dict[str, list[tuple[str, list]]]:
    """{file name: [(sheet name, [row tuples incl. header])]} in file order; no temp files allowed.
    The Timestamp column is the save time, so it is compared for presence/ISO format only."""
    assert not list(d.glob("*.tmp*")), "leftover temp files"
    out: dict[str, list[tuple[str, list]]] = {}
    for path in sorted(d.glob("*.xlsx")):
        book = openpyxl.load_workbook(path)
        sheets = []
        for ws in book.worksheets:
            rows = []
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                cells = list(row)
                if normalize_timestamps and i > 0 and len(cells) > 3 and cells[3] is not None:
                    # a save writes an ISO time (compared by format only); the placeholders the
                    # scenarios seed rows with ("t0", "t-old", ...) are compared exactly as they are
                    if isinstance(cells[3], str) and _ISO_RE.match(cells[3]):
                        cells[3] = "<iso>"
                rows.append(tuple(cells))
            sheets.append((ws.title, rows))
        out[path.name] = sheets
    return out


def data_rows(d: Path, user_id: ObjectId, name: str) -> list[tuple[str, tuple]]:
    book = openpyxl.load_workbook(xlsx_path(d, user_id, name))
    return [(ws.title, tuple(r)) for ws in book for r in ws.iter_rows(min_row=2, values_only=True)]


async def snapshot(db: Any, uids: list[ObjectId]) -> dict[str, list[dict]]:
    return {c: await db[c].find({"userId": {"$in": uids}}).to_list(None) for c in COLLECTIONS}


async def restore(db: Any, uids: list[ObjectId], snap: dict[str, list[dict]]) -> None:
    for c in COLLECTIONS:
        await db[c].delete_many({"userId": {"$in": uids}})
        if snap[c]:
            await db[c].insert_many([dict(x) for x in snap[c]])


async def capture(db: Any, uids: list[ObjectId], d: Path) -> dict[str, Any]:
    """Normalized final state: DB collections (ids/times of NEW records ignored) + the files."""
    snap = await snapshot(db, uids)
    wb_by_id = {w["_id"]: w["filename"] for w in snap["workbooks"]}
    drop = {"_id", "exportedAt", "createdAt", "updatedAt"}
    exported = []
    for rec in snap["exportedrows"]:
        item = {k: v for k, v in rec.items() if k not in drop}
        item["workbookId"] = wb_by_id.get(rec.get("workbookId"), str(rec.get("workbookId")))
        item["documentId"] = str(rec["documentId"])
        item["userId"] = str(rec["userId"])
        exported.append(item)
    exported.sort(key=lambda r: repr(sorted(r.items())))
    workbooks = sorted(
        (
            {
                **{k: v for k, v in w.items() if k not in {"_id", "createdAt", "updatedAt"}},
                "archivedAt": bool(w.get("archivedAt")),
            }
            for w in snap["workbooks"]
        ),
        key=lambda w: repr(sorted(w.items())),
    )
    settings = sorted(
        ({k: v for k, v in s.items() if k not in {"_id", "updatedAt"}} for s in snap["settings"]),
        key=lambda s: repr(sorted(s.items())),
    )
    documents = sorted(
        ({k: v for k, v in x.items() if k != "updatedAt"} for x in snap["documents"]),
        key=lambda x: str(x["_id"]),
    )
    audit = sorted(
        (
            (a["action"], tuple(sorted(a["context"].items())))
            for a in snap["auditlogs"]
            if a["action"] == "document_exported"
        ),
        key=repr,
    )
    return {
        "exportedrows": exported,
        "workbooks": workbooks,
        "settings": settings,
        "documents": documents,
        "audit": audit,
        "files": read_dir(d),
    }


async def seed_book(db: Any, user_id: ObjectId, name: str, year: int, active: bool) -> ObjectId:
    now = datetime.now(UTC)
    res = await db.workbooks.insert_one(
        {
            "userId": user_id,
            "year": year,
            "filename": name,
            "isActive": active,
            "archivedAt": None if active else now,
            "createdAt": now,
            "updatedAt": now,
        }
    )
    return res.inserted_id


async def seed_settings(db: Any, user_id: ObjectId, name: str, year: int) -> None:
    await db.settings.insert_one(
        {
            "userId": user_id,
            "key": "excelState",
            "activeWorkbookName": name,
            "activeYear": year,
            "updatedAt": datetime.now(UTC),
        }
    )


async def seed_old_record(
    db: Any,
    user_id: ObjectId,
    doc_id: ObjectId,
    workbook_id: ObjectId,
    workbook_name: str,
    sheet: str,
    number: str,
    date: str,
    when: datetime | None = None,
) -> ObjectId:
    """An earlier export (Delivery Challan) of `doc_id`; the caller writes the matching row."""
    stamp = when or datetime.now(UTC) - timedelta(days=1)
    res = await db.exportedrows.insert_one(
        {
            "documentId": doc_id,
            "userId": user_id,
            "workbookId": workbook_id,
            "workbookFilename": workbook_name,
            "worksheet": sheet,
            "documentType": "Delivery Challan",
            "taxInvoiceNo": None,
            "referenceNo": None,
            "number": number,
            "date": date,
            "exportedAt": stamp,
            "createdAt": stamp,
            "updatedAt": stamp,
        }
    )
    return res.inserted_id


async def post_bulk(client: Any, user: Any, ids: list[Any]) -> Any:
    return await client.post(
        f"{BASE}/bulk-save", headers=user.headers, json={"documentIds": [str(i) for i in ids]}
    )
