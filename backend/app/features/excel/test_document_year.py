"""The document's OWN date decides both the workbook year and the sheet month; unusable or
implausible dates fall back to the save date (IST). Missing past-year workbooks are created
on demand (archived, never active); the current year without a workbook still 409s."""

import asyncio
from datetime import UTC, datetime, timedelta, timezone
from pathlib import Path
from typing import Any

import openpyxl
import pytest
from bson import ObjectId

from app.features.excel import service as excel_service

IST = timezone(timedelta(hours=5, minutes=30))
BASE = "/api/documents"
Clock = list[datetime]


@pytest.fixture
def exports_dir(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    monkeypatch.setattr(excel_service, "EXPORT_DIR", tmp_path)
    return tmp_path


@pytest.fixture
def clock(monkeypatch: pytest.MonkeyPatch) -> Clock:
    """Mutable fake IST clock: clock[0] = datetime(...) moves 'now'."""
    now: Clock = [datetime(2026, 1, 5, 10, 0, tzinfo=IST)]
    monkeypatch.setattr(excel_service, "now_ist", lambda: now[0])
    return now


def at(y: int, m: int, d: int) -> datetime:
    return datetime(y, m, d, 10, 0, tzinfo=IST)


async def _doc(db: Any, owner: Any, number: str, date: Any) -> ObjectId:
    now = datetime.now(UTC)
    res = await db.documents.insert_one(
        {
            "userId": owner.id,
            "autoName": number,
            "originalFilename": f"{number}.pdf",
            "mimeType": "application/pdf",
            "size": 10,
            "uploadStatus": "processed",
            "documentType": "Delivery Challan",
            "number": number,
            "date": date,
            "edited": False,
            "exported": False,
            "isDeleted": False,
            "filePurged": True,
            "createdAt": now,
            "updatedAt": now,
        }
    )
    return res.inserted_id


async def _start(client: Any, user: Any, name: str = "Book2026") -> None:
    r = await client.post(f"{BASE}/new-excel-file", headers=user.headers, json={"filename": name})
    assert r.status_code == 200, r.text


async def _save(client: Any, user: Any, doc_id: Any) -> Any:
    return await client.post(f"{BASE}/{doc_id}/save", headers=user.headers)


async def _save_ok(client: Any, user: Any, doc_id: Any) -> dict:
    r = await _save(client, user, doc_id)
    assert r.status_code == 200, r.text
    return r.json()


def _rows(exports_dir: Path, user: Any, wb: dict) -> list[tuple[str, str]]:
    """(sheet title, number) for every data row of that workbook doc."""
    name = excel_service.physical_workbook_filename(user.id, wb["filename"])
    book = openpyxl.load_workbook(exports_dir / f"{name}.xlsx")
    return [
        (ws.title, str(row[1])) for ws in book for row in ws.iter_rows(min_row=2, values_only=True)
    ]


async def _wb(db: Any, user: Any, year: int) -> dict:
    wb = await db.workbooks.find_one({"userId": user.id, "year": year})
    assert wb is not None, f"no workbook for {year}"
    return wb


async def _where(db: Any, exports_dir: Path, user: Any, number: str) -> list[tuple[int, str]]:
    """Every (workbook year, sheet) in which `number` appears."""
    found = []
    async for wb in db.workbooks.find({"userId": user.id}):
        found += [(wb["year"], s) for s, n in _rows(exports_dir, user, wb) if n == number]
    return found


async def _state(db: Any, user: Any) -> dict:
    return await db.settings.find_one({"userId": user.id, "key": "excelState"})


# 1
async def test_paper_from_last_year_goes_to_that_years_workbook(
    client: Any, db: Any, make_user: Any, exports_dir: Path, clock: Clock
) -> None:
    u = await make_user()
    await _start(client, u)
    d = await _doc(db, u, "820000101", "20/12/2025")
    body = await _save_ok(client, u, d)
    assert body["worksheet"] == "December" and body["dateFallback"] is False
    assert await _where(db, exports_dir, u, "820000101") == [(2025, "December")]
    wb26 = await _wb(db, u, 2026)
    assert all(n != "820000101" for _, n in _rows(exports_dir, u, wb26))
    st = await _state(db, u)
    assert st["activeYear"] == 2026 and st["activeWorkbookName"] == "Book2026"


# 2
async def test_paper_in_current_year_unchanged(
    client: Any, db: Any, make_user: Any, exports_dir: Path, clock: Clock
) -> None:
    u = await make_user()
    clock[0] = at(2026, 1, 15)
    await _start(client, u)
    d = await _doc(db, u, "820000201", "15/01/2026")
    body = await _save_ok(client, u, d)
    assert body["worksheet"] == "January" and body["dateFallback"] is False
    assert await _where(db, exports_dir, u, "820000201") == [(2026, "January")]
    assert await db.workbooks.count_documents({"userId": u.id}) == 1


# 3
async def test_new_year_boundary(
    client: Any, db: Any, make_user: Any, exports_dir: Path, clock: Clock
) -> None:
    u = await make_user()
    clock[0] = at(2026, 1, 2)
    await _start(client, u)
    a = await _doc(db, u, "820000301", "31/12/2025")
    b = await _doc(db, u, "820000302", "01/01/2026")
    await _save_ok(client, u, a)
    await _save_ok(client, u, b)
    assert await _where(db, exports_dir, u, "820000301") == [(2025, "December")]
    assert await _where(db, exports_dir, u, "820000302") == [(2026, "January")]


# 4
@pytest.mark.parametrize("bad", [None, "", "abc", "32/13/2025", "31/02/2026", "2025-12-20"])
async def test_unusable_dates_fall_back_to_save_date(
    client: Any, db: Any, make_user: Any, exports_dir: Path, clock: Clock, bad: Any
) -> None:
    u = await make_user()
    clock[0] = at(2026, 3, 10)
    await _start(client, u)
    d = await _doc(db, u, "820000401", bad)
    body = await _save_ok(client, u, d)
    assert body["dateFallback"] is True and body["worksheet"] == "March"
    assert await _where(db, exports_dir, u, "820000401") == [(2026, "March")]


# 5
@pytest.mark.parametrize("bad", ["01/01/2052", "01/01/2005", "12/03/2026"])
async def test_implausible_dates_fall_back_and_create_no_workbook(
    client: Any, db: Any, make_user: Any, exports_dir: Path, clock: Clock, bad: str
) -> None:
    u = await make_user()
    clock[0] = at(2026, 3, 10)
    await _start(client, u)
    d = await _doc(db, u, "820000501", bad)  # 12/03/2026 = today + 2 days
    body = await _save_ok(client, u, d)
    assert body["dateFallback"] is True and body["worksheet"] == "March"
    assert await _where(db, exports_dir, u, "820000501") == [(2026, "March")]
    assert await db.workbooks.count_documents({"userId": u.id}) == 1
    assert len(list(exports_dir.glob("*.xlsx"))) == 1


def test_constants_exist() -> None:
    assert excel_service.MAX_FUTURE_DAYS == 1 and excel_service.MAX_PAST_YEARS == 5


@pytest.mark.parametrize(
    "today, date, plausible",
    [
        (at(2026, 3, 10), "11/03/2026", True),  # today + 1
        (at(2026, 3, 10), "12/03/2026", False),  # today + 2
        (at(2026, 3, 10), "10/03/2021", True),  # exactly 5y back
        (at(2026, 3, 10), "09/03/2021", False),  # one day older
        (at(2028, 2, 29), "28/02/2023", True),  # 29 Feb -> 28 Feb
        (at(2028, 2, 29), "27/02/2023", False),
        (at(2026, 12, 31), "01/01/2027", False),  # next year, inside +1 day: still falls back
    ],
)
async def test_plausibility_limits(
    client: Any,
    db: Any,
    make_user: Any,
    exports_dir: Path,
    clock: Clock,
    today: datetime,
    date: str,
    plausible: bool,
) -> None:
    u = await make_user()
    clock[0] = today
    await _start(client, u)
    d = await _doc(db, u, "820000551", date)
    r = await _save(client, u, d)
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["dateFallback"] is (not plausible)
    if plausible:
        d_, m_, y_ = (int(x) for x in date.split("/"))
        assert await _where(db, exports_dir, u, "820000551") == [(y_, excel_service.MONTHS[m_ - 1])]
    else:
        assert await _where(db, exports_dir, u, "820000551") == [
            (today.year, excel_service.MONTHS[today.month - 1])
        ]
        assert await db.workbooks.count_documents({"userId": u.id}) == 1


# 6
async def test_missing_year_workbook_created_on_demand_archived(
    client: Any, db: Any, make_user: Any, exports_dir: Path, clock: Clock
) -> None:
    u = await make_user()
    await _start(client, u)
    assert await db.workbooks.find_one({"userId": u.id, "year": 2025}) is None
    d = await _doc(db, u, "820000601", "20/12/2025")
    await _save_ok(client, u, d)
    wb = await _wb(db, u, 2025)
    assert wb and wb["isActive"] is False and wb["archivedAt"] is not None
    name = excel_service.physical_workbook_filename(u.id, wb["filename"])
    book = openpyxl.load_workbook(exports_dir / f"{name}.xlsx")
    assert book.sheetnames == ["December"]
    assert [c.value for c in book["December"][1]] == [
        "Document Type",
        "Number",
        "Date",
        "Timestamp",
    ]
    st = await _state(db, u)
    assert st["activeYear"] == 2026 and st["activeWorkbookName"] == "Book2026"
    assert (await _wb(db, u, 2026))["isActive"] is True


# 7
async def test_two_papers_back_to_back_each_in_own_workbook(
    client: Any, db: Any, make_user: Any, exports_dir: Path, clock: Clock
) -> None:
    u = await make_user()
    await _start(client, u)
    a = await _doc(db, u, "820000701", "20/12/2025")
    b = await _doc(db, u, "820000702", "04/01/2026")
    await _save_ok(client, u, a)
    await _save_ok(client, u, b)
    assert await _where(db, exports_dir, u, "820000701") == [(2025, "December")]
    assert await _where(db, exports_dir, u, "820000702") == [(2026, "January")]


# 8
async def test_bulk_save_mixed_years(
    client: Any, db: Any, make_user: Any, exports_dir: Path, clock: Clock
) -> None:
    u = await make_user()
    await _start(client, u)
    clock[0] = at(2026, 1, 5)
    dates = ["05/03/2024", "06/04/2024", "20/12/2025", "21/11/2025", "22/10/2025"]
    dates += ["01/01/2026", "02/01/2026", "03/01/2026", "04/01/2026", "05/01/2026"]
    ids = [await _doc(db, u, f"8200008{i:02d}", dt) for i, dt in enumerate(dates)]
    r = await client.post(
        f"{BASE}/bulk-save", headers=u.headers, json={"documentIds": [str(i) for i in ids]}
    )
    assert r.status_code == 200, r.text
    assert len(r.json()["succeeded"]) == 10 and r.json()["blocked"] is None
    for y in (2024, 2025, 2026):
        assert await db.workbooks.count_documents({"userId": u.id, "year": y}) == 1
    for i, dt in enumerate(dates):
        _, m, y = (int(x) for x in dt.split("/"))
        assert await _where(db, exports_dir, u, f"8200008{i:02d}") == [
            (y, excel_service.MONTHS[m - 1])
        ]
    assert len(list(exports_dir.glob("*.xlsx"))) == 3


# 9
async def test_concurrent_saves_create_one_workbook(
    client: Any, db: Any, make_user: Any, exports_dir: Path, clock: Clock
) -> None:
    u = await make_user()
    await _start(client, u)
    a = await _doc(db, u, "820000901", "20/12/2025")
    b = await _doc(db, u, "820000902", "21/12/2025")
    ra, rb = await asyncio.gather(_save(client, u, a), _save(client, u, b))
    assert ra.status_code == 200 and rb.status_code == 200, (ra.text, rb.text)
    assert await db.workbooks.count_documents({"userId": u.id, "year": 2025}) == 1
    wb = await _wb(db, u, 2025)
    assert sorted(_rows(exports_dir, u, wb)) == [
        ("December", "820000901"),
        ("December", "820000902"),
    ]


# 10
async def test_editing_date_after_export_only_flips_exported(
    client: Any, db: Any, make_user: Any, exports_dir: Path, clock: Clock
) -> None:
    u = await make_user()
    await _start(client, u)
    d = await _doc(db, u, "820001001", "20/12/2025")
    await _save_ok(client, u, d)
    r = await client.patch(
        f"{BASE}/{d}/correct", headers=u.headers, json={"field": "date", "value": "15/01/2026"}
    )
    assert r.status_code == 200, r.text
    assert (await db.documents.find_one({"_id": d}))["exported"] is False
    assert await _where(db, exports_dir, u, "820001001") == [(2025, "December")]
    assert await db.workbooks.count_documents({"userId": u.id}) == 2


# 11
async def test_fallback_row_is_removable(
    client: Any, db: Any, make_user: Any, exports_dir: Path, clock: Clock
) -> None:
    u = await make_user()
    clock[0] = at(2026, 3, 10)
    await _start(client, u)
    d = await _doc(db, u, "820001101", "garbage")
    await _save_ok(client, u, d)
    wb = await _wb(db, u, 2026)
    assert _rows(exports_dir, u, wb) == [("March", "820001101")]
    er = await db.exportedrows.find_one({"documentId": d})
    name = excel_service.physical_workbook_filename(u.id, wb["filename"])
    removed, now_empty = await excel_service.remove_rows(name, [er])
    assert removed == 1
    # a fully emptied workbook may be deleted by remove_rows itself
    assert now_empty or _rows(exports_dir, u, wb) == []


# current year without a workbook keeps the old behaviour
async def test_current_year_without_workbook_still_409(
    client: Any, db: Any, make_user: Any, exports_dir: Path, clock: Clock
) -> None:
    u = await make_user()
    await _start(client, u)
    await db.workbooks.delete_many({"userId": u.id})
    await db.settings.update_one(
        {"userId": u.id, "key": "excelState"}, {"$set": {"activeYear": 2025}}
    )
    d = await _doc(db, u, "820001201", "04/01/2026")
    r = await _save(client, u, d)
    assert r.status_code == 409
    assert "NEED_NEW_WORKBOOK" in r.text
    assert await db.workbooks.count_documents({"userId": u.id}) == 0
