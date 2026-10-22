"""ORACLE for the batch-save differential tests: a frozen, verbatim copy of the OLD per-document
Excel save (git commit 7a965d1: excel/service.py file I/O + excel/router.py save/bulk loop).

Nothing here is imported by the app. The tests run this code against copy A of a scenario and the
real HTTP API against copy B, then compare. Only adaptations from the original:
  * the router functions take a plain `user_id` instead of a CurrentUser;
  * the file I/O helpers are copied here (so later service refactors cannot change the oracle) and
    resolve paths through `excel_service.file_path`, i.e. the tests' monkeypatched EXPORT_DIR;
  * the module has its own lock dicts (the app's locks must not be shared with the oracle).
Constants/helpers that are NOT changing are taken from the current service: MONTHS,
document_period, month_from_date, now_ist, current_period, FileLockedError,
physical_workbook_filename, file_path, HEADERS.
"""

import asyncio
import os
from collections import Counter, defaultdict
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

import openpyxl
from bson import ObjectId
from fastapi import HTTPException
from filelock import FileLock
from loguru import logger
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.worksheet.worksheet import Worksheet

from app.core.database import get_database
from app.features.excel import service as excel_service

MAX_BULK_SAVE = 100
BLOCKING_ERRORS = {"NO_ACTIVE_WORKBOOK", "NEED_NEW_WORKBOOK"}
FILE_BUSY_MESSAGE = "The Excel file is busy or open. Close it in Excel and try again."

_write_locks: dict[str, asyncio.Lock] = {}
_period_locks: dict[tuple[ObjectId, int], asyncio.Lock] = {}
_doc_locks: dict[ObjectId, asyncio.Lock] = {}

_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


# ---------------------------------------------------------------- old service.py (file I/O)


def _lock_path(target: Path) -> Path:
    return target.with_suffix(".lock")


def _add_header_row(sheet: Worksheet) -> None:
    sheet.append(excel_service.HEADERS)
    for cell in sheet[1]:
        cell.font = openpyxl.styles.Font(bold=True)


def _create_workbook_sync(filename: str, month: str) -> Path:
    excel_service.EXPORT_DIR.mkdir(parents=True, exist_ok=True)
    workbook = openpyxl.Workbook()
    default_sheet = workbook.active
    workbook.remove(default_sheet)
    sheet = workbook.create_sheet(month)
    _add_header_row(sheet)
    target = excel_service.file_path(filename)
    tmp_target = target.with_suffix(f".tmp{os.getpid()}.xlsx")
    try:
        workbook.save(tmp_target)
        os.replace(tmp_target, target)
    except PermissionError as exc:
        tmp_target.unlink(missing_ok=True)
        raise excel_service.FileLockedError(
            f'"{target.name}" is open in Excel. Close it and try saving again.'
        ) from exc
    except Exception:
        tmp_target.unlink(missing_ok=True)
        raise
    return target


def _create_workbook_locked_sync(filename: str, month: str) -> Path:
    excel_service.EXPORT_DIR.mkdir(parents=True, exist_ok=True)
    target = excel_service.file_path(filename)
    with FileLock(str(_lock_path(target)), timeout=30):
        return _create_workbook_sync(filename, month)


def _get_write_lock(filename: str) -> asyncio.Lock:
    return _write_locks.setdefault(filename, asyncio.Lock())


async def create_workbook(filename: str, month: str | None = None) -> Path:
    month = month or excel_service.current_period()[1]
    async with _get_write_lock(filename):
        return await asyncio.to_thread(_create_workbook_locked_sync, filename, month)


def _clean_cell(value: object) -> object:
    if isinstance(value, str):
        return ILLEGAL_CHARACTERS_RE.sub("", value)
    return value


def _force_text_cells(cells: tuple) -> None:
    for cell in cells:
        if isinstance(cell.value, str) and cell.value.startswith(_FORMULA_PREFIXES):
            cell.data_type = "s"


def _format_number_cell(row: dict) -> str | None:
    if row["documentType"] == "Tax Invoice":
        parts = [p for p in (row.get("taxInvoiceNo"), row.get("referenceNo")) if p]
        return " / ".join(parts) if parts else None
    return row.get("number") or None


def _append_row_sync(filename: str, month: str, row: dict) -> Path:
    excel_service.EXPORT_DIR.mkdir(parents=True, exist_ok=True)
    target = excel_service.file_path(filename)
    with FileLock(str(_lock_path(target)), timeout=30):
        if not target.exists():
            _create_workbook_sync(filename, month)
        tmp_target = target.with_suffix(f".tmp{os.getpid()}.xlsx")
        try:
            workbook = openpyxl.load_workbook(target)
            if month in workbook.sheetnames:
                sheet = workbook[month]
            else:
                sheet = workbook.create_sheet(month)
                _add_header_row(sheet)
            sheet.append(
                [
                    _clean_cell(row["documentType"]),
                    _clean_cell(_format_number_cell(row)),
                    _clean_cell(row["date"]),
                    _clean_cell(row["timestamp"]),
                ]
            )
            _force_text_cells(sheet[sheet.max_row])
            workbook.save(tmp_target)
            os.replace(tmp_target, target)
        except PermissionError as exc:
            tmp_target.unlink(missing_ok=True)
            raise excel_service.FileLockedError(
                f'"{target.name}" is open in Excel. Close it and try saving again.'
            ) from exc
        except Exception:
            tmp_target.unlink(missing_ok=True)
            raise
        return target


async def append_row(filename: str, month: str, row: dict) -> Path:
    async with _get_write_lock(filename):
        return await asyncio.to_thread(_append_row_sync, filename, month, row)


def row_key(row: dict) -> tuple[object, object, object]:
    return (
        _clean_cell(row["documentType"]),
        _clean_cell(_format_number_cell(row)),
        _clean_cell(row.get("date")),
    )


def _remove_rows_sync(
    filename: str,
    rows_to_remove: list[dict],
    keep_file: bool = False,
    protect: tuple[str, tuple] | None = None,
) -> tuple[int, bool]:
    target = excel_service.file_path(filename)
    if not target.exists() or not rows_to_remove:
        return 0, False

    by_sheet: dict[str, Counter] = defaultdict(Counter)
    for row in rows_to_remove:
        sheet_name = row.get("worksheet") or excel_service.month_from_date(row.get("date"))
        by_sheet[sheet_name][row_key(row)] += 1

    removed = 0
    fully_empty = False
    with FileLock(str(_lock_path(target)), timeout=30):
        if not target.exists():
            return 0, False
        workbook = openpyxl.load_workbook(target)
        for sheet_name, counter in by_sheet.items():
            if sheet_name not in workbook.sheetnames:
                continue
            sheet = workbook[sheet_name]
            remaining = Counter(counter)
            rows_to_delete: list[int] = []
            protect_key = protect[1] if protect and protect[0] == sheet_name else None
            last_protected_row = None
            for row_cells in sheet.iter_rows(min_row=2):
                key = (row_cells[0].value, row_cells[1].value, row_cells[2].value)
                if key == protect_key:
                    last_protected_row = row_cells[0].row
                if remaining.get(key, 0) > 0:
                    rows_to_delete.append(row_cells[0].row)
                    remaining[key] -= 1
            if last_protected_row in rows_to_delete:
                rows_to_delete.remove(last_protected_row)
            for idx in sorted(rows_to_delete, reverse=True):
                sheet.delete_rows(idx)
            removed += len(rows_to_delete)

        fully_empty = all(sheet.max_row <= 1 for sheet in workbook.worksheets)
        if not fully_empty or keep_file:
            tmp_target = target.with_suffix(f".tmp{os.getpid()}.xlsx")
            try:
                workbook.save(tmp_target)
                os.replace(tmp_target, target)
            except PermissionError as exc:
                tmp_target.unlink(missing_ok=True)
                raise excel_service.FileLockedError(
                    f'"{target.name}" is open in Excel. Close it and try saving again.'
                ) from exc
            except Exception:
                tmp_target.unlink(missing_ok=True)
                raise

    if fully_empty and not keep_file:
        target.unlink(missing_ok=True)
        _lock_path(target).unlink(missing_ok=True)

    return removed, fully_empty


async def remove_rows(
    filename: str,
    rows_to_remove: list[dict],
    keep_file: bool = False,
    protect: tuple[str, tuple] | None = None,
) -> tuple[int, bool]:
    async with _get_write_lock(filename):
        return await asyncio.to_thread(
            _remove_rows_sync, filename, rows_to_remove, keep_file, protect
        )


# ---------------------------------------------------------------- old router.py (save path)


async def _get_settings(user_id: ObjectId) -> dict | None:
    db = get_database()
    return await db.settings.find_one({"userId": user_id, "key": "excelState"})


async def _workbook_for_year(db: Any, user_id: ObjectId, year: int) -> dict | None:
    return await db.workbooks.find_one(
        {"userId": user_id, "year": year}, sort=[("isActive", -1), ("createdAt", -1)]
    )


async def _create_year_workbook(db: Any, user_id: ObjectId, year: int, month: str) -> dict:
    async with _period_locks.setdefault((user_id, year), asyncio.Lock()):
        existing = await _workbook_for_year(db, user_id, year)
        if existing:
            return existing
        taken = {w["filename"] async for w in db.workbooks.find({"userId": user_id})}
        name, n = f"Excel {year}", 1
        while name in taken:
            n += 1
            name = f"Excel {year} ({n})"
        await create_workbook(excel_service.physical_workbook_filename(user_id, name), month)
        now = datetime.now(UTC)
        record = {
            "userId": user_id,
            "year": year,
            "filename": name,
            "isActive": False,
            "archivedAt": now,
            "createdAt": now,
            "updatedAt": now,
        }
        record["_id"] = (await db.workbooks.insert_one(record)).inserted_id
        return record


async def _replace_old_rows(
    db: Any,
    user_id: ObjectId,
    doc_id: ObjectId,
    new_record_id: ObjectId,
    new_file: str,
    new_sheet: str,
    new_key: tuple,
    new_file_existed: bool,
) -> tuple[int, str | None]:
    old = (
        await db.exportedrows.find(
            {"documentId": doc_id, "userId": user_id, "_id": {"$ne": new_record_id}}
        )
        .sort("exportedAt", 1)
        .to_list(None)
    )
    by_workbook: dict[Any, list[dict]] = defaultdict(list)
    for rec in old:
        by_workbook[(rec.get("workbookId"), rec.get("workbookFilename"))].append(rec)

    replaced = 0
    warnings: list[str] = []
    for (workbook_id, workbook_filename), recs in by_workbook.items():
        record_ids = [r["_id"] for r in recs]
        workbook = (
            await db.workbooks.find_one({"_id": workbook_id, "userId": user_id})
            if workbook_id
            else None
        )
        if not workbook and workbook_id is None and workbook_filename:
            workbook = {"filename": workbook_filename}
        if not workbook:
            await db.exportedrows.delete_many({"_id": {"$in": record_ids}})
            continue
        physical = excel_service.physical_workbook_filename(user_id, workbook["filename"])
        if not excel_service.file_path(physical).exists() or (
            physical == new_file and not new_file_existed
        ):
            await db.exportedrows.delete_many({"_id": {"$in": record_ids}})
            continue
        name = workbook["filename"]
        try:
            protect = (new_sheet, new_key) if physical == new_file else None
            removed, _ = await remove_rows(physical, recs, keep_file=True, protect=protect)
        except excel_service.FileLockedError:
            warnings.append(
                f"The new row was saved, but the old row could not be removed from {name} "
                "because it is open in Excel. Close it and press Save Again to finish."
            )
            continue
        except Exception as exc:  # noqa: BLE001
            logger.error(
                f"Removing the old Excel row of document {doc_id} failed: "
                f"class={type(exc).__name__}"
            )
            warnings.append(
                f"The new row was saved, but the old row could not be removed from {name}. "
                "Press Save Again to try again."
            )
            continue
        replaced += removed
        await db.exportedrows.delete_many({"_id": {"$in": record_ids}})
        if removed < len(recs):
            sheet = recs[0].get("worksheet") or excel_service.month_from_date(recs[0].get("date"))
            warnings.append(
                f"The old row for this document was not found in {name} ({sheet}). "
                "Please check that file by hand."
            )
    return replaced, (" ".join(warnings) or None)


async def save_document_to_excel(doc_id: ObjectId, user_id: ObjectId) -> dict:
    async with _doc_locks.setdefault(doc_id, asyncio.Lock()):
        return await _save_document_locked(doc_id, user_id)


async def _save_document_locked(doc_id: ObjectId, user_id: ObjectId) -> dict:
    db = get_database()
    doc = await db.documents.find_one(
        {"_id": doc_id, "userId": user_id, "isDeleted": {"$ne": True}}
    )
    if not doc:
        raise HTTPException(status_code=404, detail="Document not found.")
    if doc["uploadStatus"] != "processed":
        raise HTTPException(status_code=400, detail="Document has not been processed yet.")

    now_local = excel_service.now_ist()
    current_year, _ = excel_service.current_period(now_local)
    year, sheet_month, date_used = excel_service.document_period(doc.get("date"), now_local)
    settings = await _get_settings(user_id)
    if not settings or not settings.get("activeWorkbookName"):
        raise HTTPException(
            status_code=400,
            detail={
                "error": "NO_ACTIVE_WORKBOOK",
                "year": current_year,
                "message": "No active Excel workbook yet. Name your first workbook to continue.",
            },
        )

    now = datetime.now(UTC)
    workbook_doc = await _workbook_for_year(db, user_id, year)
    if workbook_doc is None and year == settings.get("activeYear"):
        workbook_doc = {"_id": None, "filename": settings["activeWorkbookName"], "year": year}
    if workbook_doc is None and year == current_year:
        if settings.get("activeYear") != year:
            await db.workbooks.update_many(
                {"userId": user_id, "isActive": True, "year": settings.get("activeYear")},
                {"$set": {"isActive": False, "archivedAt": now, "updatedAt": now}},
            )
        raise HTTPException(
            status_code=409,
            detail={
                "error": "NEED_NEW_WORKBOOK",
                "year": year,
                "message": (
                    f"The year changed to {year}. Create a new workbook for {year} to continue."
                ),
            },
        )
    if workbook_doc is None:
        try:
            workbook_doc = await _create_year_workbook(db, user_id, year, sheet_month)
        except excel_service.FileLockedError as exc:
            raise HTTPException(status_code=409, detail=FILE_BUSY_MESSAGE) from exc

    row = {
        "documentType": doc["documentType"],
        "taxInvoiceNo": doc.get("taxInvoiceNo"),
        "referenceNo": doc.get("referenceNo"),
        "number": doc.get("number"),
        "date": doc.get("date"),
        "timestamp": now.isoformat(),
    }

    active_filename = workbook_doc["filename"]
    new_file_existed = excel_service.file_path(
        excel_service.physical_workbook_filename(user_id, active_filename)
    ).exists()
    try:
        await append_row(
            excel_service.physical_workbook_filename(user_id, active_filename), sheet_month, row
        )
    except excel_service.FileLockedError as exc:
        raise HTTPException(status_code=409, detail=FILE_BUSY_MESSAGE) from exc

    new_record = await db.exportedrows.insert_one(
        {
            "documentId": doc["_id"],
            "userId": user_id,
            "workbookId": workbook_doc["_id"],
            "workbookFilename": active_filename,
            "worksheet": sheet_month,
            "documentType": row["documentType"],
            "taxInvoiceNo": row["taxInvoiceNo"],
            "referenceNo": row["referenceNo"],
            "number": row["number"],
            "date": row["date"],
            "exportedAt": now,
            "createdAt": now,
            "updatedAt": now,
        }
    )
    try:
        replaced, stale_warning = await _replace_old_rows(
            db,
            user_id,
            doc["_id"],
            new_record.inserted_id,
            excel_service.physical_workbook_filename(user_id, active_filename),
            sheet_month,
            row_key(row),
            new_file_existed,
        )
    except Exception as exc:  # noqa: BLE001
        logger.error(
            f"Replacing the old Excel rows of document {doc_id} failed: class={type(exc).__name__}"
        )
        replaced = 0
        stale_warning = (
            "The new row was saved, but the old row could not be checked. "
            "Press Save Again to try again."
        )
    await db.documents.update_one({"_id": doc["_id"]}, {"$set": {"exported": True}})

    from app.core.audit_log import log_action

    await log_action(
        user_id,
        "document_exported",
        {
            "documentId": str(doc["_id"]),
            "worksheet": sheet_month,
            "workbook": active_filename,
            "replaced": replaced,
        },
    )
    result: dict = {
        "message": "Excel file appended successfully.",
        "worksheet": sheet_month,
        "workbook": active_filename,
        "year": year,
        "dateFallback": not date_used,
        "replaced": replaced,
    }
    if stale_warning:
        result["staleRowWarning"] = stale_warning
    return result


async def single_save(doc_id: ObjectId, user_id: ObjectId) -> tuple[int, dict]:
    """The old POST /{doc_id}/save as (status code, JSON body)."""
    try:
        return 200, await save_document_to_excel(doc_id, user_id)
    except HTTPException as exc:
        return exc.status_code, {"detail": exc.detail}


async def bulk_save(raw_ids: list[ObjectId], user_id: ObjectId) -> tuple[int, dict]:
    """The old POST /bulk-save as (status code, JSON body): the sequential loop, verbatim."""
    ids = list(dict.fromkeys(raw_ids))
    if len(ids) > MAX_BULK_SAVE:
        return 400, {"detail": f"You can save at most {MAX_BULK_SAVE} documents at once."}

    succeeded: list[str] = []
    failed: list[dict] = []
    date_fallback: list[str] = []
    blocked: dict | None = None
    not_attempted: list[str] = []
    replaced_total = 0
    stale_warnings: list[dict] = []
    for index, doc_id in enumerate(ids):
        try:
            result = await save_document_to_excel(doc_id, user_id)
            succeeded.append(str(doc_id))
            replaced_total += result.get("replaced", 0)
            if result.get("staleRowWarning"):
                stale_warnings.append(
                    {"documentId": str(doc_id), "message": result["staleRowWarning"]}
                )
            if result.get("dateFallback"):
                date_fallback.append(str(doc_id))
        except HTTPException as exc:
            detail = exc.detail
            if isinstance(detail, dict) and detail.get("error") in BLOCKING_ERRORS:
                blocked = {
                    "error": detail["error"],
                    "year": detail.get("year"),
                    "message": detail.get("message", "Create an Excel workbook to continue."),
                }
                not_attempted = [str(i) for i in ids[index:]]
                break
            reason = detail if isinstance(detail, str) else detail.get("message", "Could not save.")
            failed.append({"documentId": str(doc_id), "reason": reason})
        except Exception as exc:  # noqa: BLE001
            logger.error(f"Bulk save of document {doc_id} failed: class={type(exc).__name__}")
            failed.append({"documentId": str(doc_id), "reason": "Could not save this document."})

    return 200, {
        "message": f"{len(succeeded)}/{len(ids)} saved successfully.",
        "succeeded": succeeded,
        "failed": failed,
        "blocked": blocked,
        "notAttempted": not_attempted,
        "dateFallback": date_fallback,
        "replaced": replaced_total,
        "staleWarnings": stale_warnings,
    }
