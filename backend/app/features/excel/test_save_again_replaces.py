"""Save Again REPLACES: a repeat save appends the new row and then removes the document's
earlier rows (matched by the stored values) and their exportedrows records. Failures never
lose the new row; blocked saves change nothing."""

import asyncio
import os
from datetime import UTC, datetime, timedelta, timezone
from pathlib import Path
from typing import Any

import openpyxl
import pytest
from bson import ObjectId

from app.features.admin import purge_service
from app.features.excel import service as excel_service

IST = timezone(timedelta(hours=5, minutes=30))
BASE = "/api/documents"
WB_NAME = "Book2026"
Clock = list[datetime]
DATE = "05/03/2026"


@pytest.fixture
def exports_dir(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    monkeypatch.setattr(excel_service, "EXPORT_DIR", tmp_path)
    return tmp_path


@pytest.fixture
def clock(monkeypatch: pytest.MonkeyPatch) -> Clock:
    now: Clock = [datetime(2026, 3, 10, 10, 0, tzinfo=IST)]
    monkeypatch.setattr(excel_service, "now_ist", lambda: now[0])
    return now


def at(y: int, m: int, d: int) -> datetime:
    return datetime(y, m, d, 10, 0, tzinfo=IST)


async def _doc(db: Any, owner: Any, number: str, date: str | None = DATE) -> ObjectId:
    now = datetime.now(UTC)
    res = await db.documents.insert_one(
        {
            "userId": owner.id,
            "autoName": number,
            "originalFilename": f"{number}.pdf",
            "mimeType": "application/pdf",
            "size": 10,
            "uploadStatus": "processed",
            "documentType": "Delivery Challan",
            "number": number,
            "date": date,
            "edited": False,
            "exported": False,
            "isDeleted": False,
            "filePurged": True,
            "createdAt": now,
            "updatedAt": now,
        }
    )
    return res.inserted_id


async def _start(client: Any, user: Any, name: str = WB_NAME) -> None:
    r = await client.post(f"{BASE}/new-excel-file", headers=user.headers, json={"filename": name})
    assert r.status_code == 200, r.text


async def _save(client: Any, user: Any, doc_id: Any) -> Any:
    return await client.post(f"{BASE}/{doc_id}/save", headers=user.headers)


async def _save_ok(client: Any, user: Any, doc_id: Any) -> dict:
    r = await _save(client, user, doc_id)
    assert r.status_code == 200, r.text
    return r.json()


async def _bulk(client: Any, user: Any, ids: list[Any]) -> dict:
    r = await client.post(
        f"{BASE}/bulk-save", headers=user.headers, json={"documentIds": [str(i) for i in ids]}
    )
    assert r.status_code == 200, r.text
    return r.json()


async def _set_date(client: Any, user: Any, doc_id: Any, value: str) -> None:
    r = await client.patch(
        f"{BASE}/{doc_id}/correct", headers=user.headers, json={"field": "date", "value": value}
    )
    assert r.status_code == 200, r.text


def _path(exports_dir: Path, user: Any, wb: dict) -> Path:
    return exports_dir / f"{excel_service.physical_workbook_filename(user.id, wb['filename'])}.xlsx"


async def _wb(db: Any, user: Any, year: int = 2026) -> dict:
    wb = await db.workbooks.find_one({"userId": user.id, "year": year})
    assert wb is not None, f"no workbook for {year}"
    return wb


async def _rows(
    db: Any, exports_dir: Path, user: Any, year: int = 2026
) -> list[tuple[str, str, str]]:
    """(sheet, number, date) of every data row in that year's workbook file."""
    path = _path(exports_dir, user, await _wb(db, user, year))
    assert path.exists(), f"{path.name} is missing"
    book = openpyxl.load_workbook(path)
    return [
        (ws.title, str(r[1]), str(r[2]))
        for ws in book
        for r in ws.iter_rows(min_row=2, values_only=True)
    ]


async def _numbers(db: Any, exports_dir: Path, user: Any, year: int = 2026) -> list[str]:
    return [n for _, n, _ in await _rows(db, exports_dir, user, year)]


async def _records(db: Any, doc_id: Any) -> int:
    return await db.exportedrows.count_documents({"documentId": doc_id})


async def _legacy(
    db: Any,
    exports_dir: Path,
    user: Any,
    doc_id: Any,
    number: str,
    date: str,
    count: int,
    with_worksheet: bool = True,
) -> None:
    """Pre-existing state: `count` rows in the 2026 workbook and `count` matching records."""
    wb = await _wb(db, user)
    physical = excel_service.physical_workbook_filename(user.id, wb["filename"])
    month = excel_service.month_from_date(date)
    now = datetime.now(UTC)
    records = []
    for _ in range(count):
        row = {
            "documentType": "Delivery Challan",
            "taxInvoiceNo": None,
            "referenceNo": None,
            "number": number,
            "date": date,
            "timestamp": "legacy",
        }
        await excel_service.append_row(physical, month, row)
        rec: dict[str, Any] = {
            "documentId": doc_id,
            "userId": user.id,
            "workbookId": wb["_id"],
            "documentType": "Delivery Challan",
            "taxInvoiceNo": None,
            "referenceNo": None,
            "number": number,
            "date": date,
            "exportedAt": now,
            "createdAt": now,
            "updatedAt": now,
        }
        if with_worksheet:
            rec["worksheet"] = month
        records.append(rec)
    await db.exportedrows.insert_many(records)


def _fail_replace_from(monkeypatch: pytest.MonkeyPatch, nth: int) -> Any:
    """os.replace raises PermissionError from the nth call on; returns a restore() function."""
    real = os.replace
    calls = [0]

    def fake(src: Any, dst: Any, *a: Any, **k: Any) -> None:
        calls[0] += 1
        if calls[0] >= nth:
            raise PermissionError("locked by Excel")
        real(src, dst, *a, **k)

    monkeypatch.setattr(excel_service.os, "replace", fake)
    return lambda: monkeypatch.setattr(excel_service.os, "replace", real)


# 1
async def test_edited_date_moves_row_to_new_year_workbook(
    client: Any, db: Any, make_user: Any, exports_dir: Path, clock: Clock
) -> None:
    u = await make_user()
    await _start(client, u)
    d = await _doc(db, u, "820000101", "15/03/2024")
    await _save_ok(client, u, d)
    assert await _numbers(db, exports_dir, u, 2024) == ["820000101"]
    await _set_date(client, u, d, "15/03/2025")
    await _save_ok(client, u, d)
    assert await _numbers(db, exports_dir, u, 2024) == []  # file still there, header only
    assert await _numbers(db, exports_dir, u, 2025) == ["820000101"]
    assert await _records(db, d) == 1


# 2
async def test_save_again_without_edit_keeps_one_row(
    client: Any, db: Any, make_user: Any, exports_dir: Path, clock: Clock
) -> None:
    u = await make_user()
    await _start(client, u)
    d = await _doc(db, u, "820000201")
    await _save_ok(client, u, d)
    body = await _save_ok(client, u, d)
    assert body["replaced"] == 1
    assert await _numbers(db, exports_dir, u) == ["820000201"]
    assert await _records(db, d) == 1


# 3
async def test_legacy_three_old_records_collapse_to_one(
    client: Any, db: Any, make_user: Any, exports_dir: Path, clock: Clock
) -> None:
    u = await make_user()
    await _start(client, u)
    d = await _doc(db, u, "820000301")
    await _legacy(db, exports_dir, u, d, "820000301", DATE, 3, with_worksheet=False)
    assert len(await _numbers(db, exports_dir, u)) == 3
    body = await _save_ok(client, u, d)
    assert body["replaced"] == 3
    assert await _numbers(db, exports_dir, u) == ["820000301"]
    assert await _records(db, d) == 1


# 4
async def test_save_all_twice_keeps_ten_rows(
    client: Any, db: Any, make_user: Any, exports_dir: Path, clock: Clock
) -> None:
    u = await make_user()
    await _start(client, u)
    ids = [await _doc(db, u, f"8200004{i:02d}") for i in range(10)]
    first = await _bulk(client, u, ids)
    assert len(first["succeeded"]) == 10
    second = await _bulk(client, u, ids)
    assert len(second["succeeded"]) == 10 and second["replaced"] == 10
    assert sorted(await _numbers(db, exports_dir, u)) == sorted(
        f"8200004{i:02d}" for i in range(10)
    )
    assert await db.exportedrows.count_documents({"userId": u.id}) == 10


# 5
async def test_locked_file_keeps_old_row_warns_then_retry_cleans_up(
    client: Any,
    db: Any,
    make_user: Any,
    exports_dir: Path,
    clock: Clock,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    u = await make_user()
    await _start(client, u)
    d = await _doc(db, u, "820000501")
    await _save_ok(client, u, d)
    restore = _fail_replace_from(monkeypatch, 2)  # append is call 1, the removal call 2
    body = await _save_ok(client, u, d)
    assert WB_NAME in body["staleRowWarning"] and "Save Again" in body["staleRowWarning"]
    assert await _numbers(db, exports_dir, u) == ["820000501", "820000501"]
    assert await _records(db, d) == 2
    restore()
    body = await _save_ok(client, u, d)
    assert body["replaced"] == 2 and "staleRowWarning" not in body
    assert await _numbers(db, exports_dir, u) == ["820000501"]
    assert await _records(db, d) == 1


# 6
async def test_old_row_deleted_by_hand_gives_not_found_warning(
    client: Any, db: Any, make_user: Any, exports_dir: Path, clock: Clock
) -> None:
    u = await make_user()
    await _start(client, u)
    d = await _doc(db, u, "820000601")
    await _save_ok(client, u, d)
    path = _path(exports_dir, u, await _wb(db, u))
    book = openpyxl.load_workbook(path)
    book["March"].delete_rows(2)
    book.save(path)
    body = await _save_ok(client, u, d)
    assert "not found" in body["staleRowWarning"] and "March" in body["staleRowWarning"]
    assert WB_NAME in body["staleRowWarning"]
    assert await _records(db, d) == 1
    assert await _numbers(db, exports_dir, u) == ["820000601"]


# 7
async def test_workbook_file_deleted_from_disk_no_warning_no_crash(
    client: Any, db: Any, make_user: Any, exports_dir: Path, clock: Clock
) -> None:
    u = await make_user()
    await _start(client, u)
    d = await _doc(db, u, "820000701")
    await _save_ok(client, u, d)
    _path(exports_dir, u, await _wb(db, u)).unlink()
    # NOTE: append recreates the missing file first, so the old row is "not found" in the
    # recreated file; the spec says no warning for a gone file (see report: ambiguity).
    body = await _save_ok(client, u, d)
    assert "staleRowWarning" not in body
    assert await _records(db, d) == 1


async def test_old_record_pointing_at_missing_workbook_is_just_dropped(
    client: Any, db: Any, make_user: Any, exports_dir: Path, clock: Clock
) -> None:
    u = await make_user()
    await _start(client, u)
    d = await _doc(db, u, "820000702")
    now = datetime.now(UTC)
    await db.exportedrows.insert_one(
        {
            "documentId": d,
            "userId": u.id,
            "workbookId": ObjectId(),  # no such workbook record
            "worksheet": "March",
            "documentType": "Delivery Challan",
            "taxInvoiceNo": None,
            "referenceNo": None,
            "number": "820000702",
            "date": DATE,
            "exportedAt": now,
            "createdAt": now,
            "updatedAt": now,
        }
    )
    body = await _save_ok(client, u, d)
    assert "staleRowWarning" not in body
    assert await _records(db, d) == 1
    assert await _numbers(db, exports_dir, u) == ["820000702"]


# 8a
async def test_rollover_409_changes_nothing(
    client: Any, db: Any, make_user: Any, exports_dir: Path, clock: Clock
) -> None:
    u = await make_user()
    await _start(client, u)
    d = await _doc(db, u, "820000801")
    await _save_ok(client, u, d)
    clock[0] = at(2027, 1, 10)
    await db.documents.update_one({"_id": d}, {"$set": {"date": "05/01/2027"}})
    r = await _save(client, u, d)
    assert r.status_code == 409 and "NEED_NEW_WORKBOOK" in r.text
    assert await _numbers(db, exports_dir, u) == ["820000801"]
    assert await _records(db, d) == 1


async def test_no_active_workbook_400_changes_nothing(
    client: Any, db: Any, make_user: Any, exports_dir: Path, clock: Clock
) -> None:
    u = await make_user()
    await _start(client, u)
    d = await _doc(db, u, "820000802")
    await _save_ok(client, u, d)
    await db.settings.delete_many({"userId": u.id})
    await db.workbooks.update_many({"userId": u.id}, {"$set": {"isActive": False}})
    r = await _save(client, u, d)
    assert r.status_code == 400 and "NO_ACTIVE_WORKBOOK" in r.text
    assert await _numbers(db, exports_dir, u) == ["820000802"]
    assert await _records(db, d) == 1


# 8b
async def test_append_failing_file_busy_changes_nothing(
    client: Any,
    db: Any,
    make_user: Any,
    exports_dir: Path,
    clock: Clock,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    u = await make_user()
    await _start(client, u)
    d = await _doc(db, u, "820000803")
    await _save_ok(client, u, d)
    _fail_replace_from(monkeypatch, 1)
    r = await _save(client, u, d)
    assert r.status_code == 409
    assert await _numbers(db, exports_dir, u) == ["820000803"]
    assert await _records(db, d) == 1


# 9
async def test_concurrent_save_again_of_same_document_ends_with_one_row(
    client: Any, db: Any, make_user: Any, exports_dir: Path, clock: Clock
) -> None:
    u = await make_user()
    await _start(client, u)
    d = await _doc(db, u, "820000901")
    await _save_ok(client, u, d)
    ra, rb = await asyncio.gather(_save(client, u, d), _save(client, u, d))
    assert ra.status_code == 200 and rb.status_code == 200, (ra.text, rb.text)
    assert await _numbers(db, exports_dir, u) == ["820000901"]
    assert await _records(db, d) == 1


# 10
async def test_two_documents_with_identical_values_keep_two_rows(
    client: Any, db: Any, make_user: Any, exports_dir: Path, clock: Clock
) -> None:
    u = await make_user()
    await _start(client, u)
    a = await _doc(db, u, "820001001")
    b = await _doc(db, u, "820001001")
    await _save_ok(client, u, a)
    await _save_ok(client, u, b)
    await _save_ok(client, u, a)
    # Documented limit: rows have no ID, so with identical type+number+date the WRONG physical
    # row (b's) may be the one removed. Only the counts are guaranteed: 2 rows, 1 record each.
    assert await _numbers(db, exports_dir, u) == ["820001001", "820001001"]
    assert await _records(db, a) == 1 and await _records(db, b) == 1
    assert await db.exportedrows.count_documents({"userId": u.id}) == 2


# 11
async def test_other_users_data_is_never_touched(
    client: Any, db: Any, make_user: Any, exports_dir: Path, clock: Clock
) -> None:
    a, b = await make_user(), await make_user()
    await _start(client, a)
    await _start(client, b)
    da = await _doc(db, a, "820001101")
    db_ = await _doc(db, b, "820001101")  # same number on purpose
    await _save_ok(client, b, db_)
    await _save_ok(client, a, da)
    before = [r async for r in db.exportedrows.find({"userId": b.id})]
    await _save_ok(client, a, da)
    assert [r async for r in db.exportedrows.find({"userId": b.id})] == before
    assert await _numbers(db, exports_dir, b) == ["820001101"]
    assert await _numbers(db, exports_dir, a) == ["820001101"]
    assert await _records(db, da) == 1


# 12
async def test_purge_after_replace_removes_the_row_and_history_lists_document_once(
    client: Any, db: Any, make_user: Any, exports_dir: Path, clock: Clock
) -> None:
    u = await make_user()
    await _start(client, u)
    d = await _doc(db, u, "820001201")
    other = await _doc(db, u, "820001202")
    await _save_ok(client, u, d)
    await _save_ok(client, u, other)
    await _save_ok(client, u, d)
    r = await client.get(f"{BASE}/export-history", headers=u.headers)
    assert r.status_code == 200, r.text
    assert [e["documentId"] for e in r.json()["exports"]].count(str(d)) == 1
    records = [x async for x in db.exportedrows.find({"documentId": d})]
    removed, _ = await purge_service.remove_exported_rows_from_workbooks(records)
    assert removed == 1
    assert await _numbers(db, exports_dir, u) == ["820001202"]


# 13
async def test_edit_across_year_boundary_moves_row(
    client: Any, db: Any, make_user: Any, exports_dir: Path, clock: Clock
) -> None:
    u = await make_user()
    clock[0] = at(2026, 1, 2)
    await _start(client, u)
    d = await _doc(db, u, "820001301", "31/12/2025")
    await _save_ok(client, u, d)
    assert await _numbers(db, exports_dir, u, 2025) == ["820001301"]
    await _set_date(client, u, d, "01/01/2026")
    await _save_ok(client, u, d)
    assert await _numbers(db, exports_dir, u, 2025) == []
    assert await _numbers(db, exports_dir, u, 2026) == ["820001301"]
    assert await _records(db, d) == 1


# 14
async def test_fifty_old_rows_removed_other_rows_untouched_and_in_order(
    client: Any, db: Any, make_user: Any, exports_dir: Path, clock: Clock
) -> None:
    u = await make_user()
    await _start(client, u)
    wb = await _wb(db, u)
    physical = excel_service.physical_workbook_filename(u.id, wb["filename"])

    async def other_row(n: str) -> None:
        await excel_service.append_row(
            physical,
            "March",
            {
                "documentType": "Delivery Challan",
                "taxInvoiceNo": None,
                "referenceNo": None,
                "number": n,
                "date": DATE,
                "timestamp": "x",
            },
        )

    before = [f"OTHB{i:02d}" for i in range(10)]
    after = [f"OTHA{i:02d}" for i in range(10)]
    for n in before:
        await other_row(n)
    d = await _doc(db, u, "820001401")
    await _legacy(db, exports_dir, u, d, "820001401", DATE, 50)
    for n in after:
        await other_row(n)
    body = await _save_ok(client, u, d)
    assert body["replaced"] == 50
    nums = await _numbers(db, exports_dir, u)
    assert [n for n in nums if n.startswith("OTH")] == before + after
    assert nums.count("820001401") == 1
    assert await _records(db, d) == 1


# 15
async def test_response_fields(
    client: Any, db: Any, make_user: Any, exports_dir: Path, clock: Clock
) -> None:
    u = await make_user()
    await _start(client, u)
    d = await _doc(db, u, "820001501")
    body = await _save_ok(client, u, d)
    assert body["replaced"] == 0 and "staleRowWarning" not in body
    assert body["worksheet"] == "March" and "message" in body
    bulk = await _bulk(client, u, [d])
    assert bulk["replaced"] == 1 and bulk["staleWarnings"] == []
    assert bulk["succeeded"] == [str(d)]
